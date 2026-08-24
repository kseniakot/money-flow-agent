import csv
import io
from collections import defaultdict

import matplotlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402


def build_report_xlsx(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Отчёт"
    ws.sheet_properties.outlinePr.summaryBelow = False

    headers = ["Категория / продукт", "Дата", "Магазин", "Кол-во", "Цена/ед", "Сумма", "Валюта"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F7D6B")

    def _total_str(items):
        totals: dict[str, float] = defaultdict(float)
        for r in items:
            totals[r["currency"]] += r["price"] or 0
        return ", ".join(f"{v:.2f} {cur}" for cur, v in totals.items())

    cats: dict[str, list] = {}
    for r in rows:
        cats.setdefault(r["category_name"], []).append(r)

    for cat in sorted(cats):
        items = cats[cat]
        ws.append([cat, "", "", "", "", _total_str(items), ""])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E2F0EC")

        products: dict[str, list] = {}
        for r in items:
            products.setdefault(r["product_name"], []).append(r)

        for prod in sorted(products):
            pitems = products[prod]
            ws.append([prod, "", "", "", "", _total_str(pitems), ""])
            ws.row_dimensions[ws.max_row].outline_level = 1
            for cell in ws[ws.max_row]:
                cell.font = Font(bold=True, italic=True)

            for r in sorted(pitems, key=lambda x: x["purchased_at"]):
                ws.append(
                    [
                        "",
                        r["purchased_at"][:16],
                        r["place"] or "",
                        f"{r['qty']:g} {r['unit']}",
                        r["unit_price"],
                        r["price"],
                        r["currency"],
                    ]
                )
                ws.row_dimensions[ws.max_row].outline_level = 2

    for col, width in zip("ABCDEFG", [28, 18, 20, 12, 10, 10, 8]):
        ws.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def expenses_csv(rows: list[dict]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(
        ["id", "дата", "категория", "продукт", "кол-во", "единица",
         "цена_за_ед", "сумма", "валюта", "место", "источник"]
    )
    for r in rows:
        writer.writerow(
            [r["id"], r["purchased_at"], r["category_name"], r["product_name"],
             r["qty"], r["unit"], r["unit_price"], r["price"], r["currency"],
             r["place"] or "", r["source"]]
        )
    return buf.getvalue().encode("utf-8-sig")


def build_charts(rows: list[dict]) -> list[tuple[str, bytes]]:
    by_cur: dict[str, dict[str, float]] = {}
    for r in rows:
        cats = by_cur.setdefault(r["currency"], defaultdict(float))
        cats[r["category_name"]] += r["price"] or 0

    charts = []
    for cur in sorted(by_cur):
        data = sorted(by_cur[cur].items(), key=lambda x: x[1])
        labels = [d[0] for d in data]
        values = [d[1] for d in data]

        fig, ax = plt.subplots(figsize=(7, max(3, 0.5 * len(labels) + 1)))
        bars = ax.barh(labels, values, color="#0f7d6b")
        ax.set_xlabel(cur)
        ax.set_title(f"Расходы по категориям ({cur})")
        ax.spines[["top", "right"]].set_visible(False)
        for bar, v in zip(bars, values):
            ax.text(bar.get_width(), bar.get_y() + bar.get_height() / 2, f" {v:.2f}", va="center")
        fig.tight_layout()

        buf = io.BytesIO()
        fig.savefig(buf, format="png", dpi=120)
        plt.close(fig)
        charts.append((cur, buf.getvalue()))
    return charts
