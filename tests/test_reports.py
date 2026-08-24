from app.bot import reports


def row(product, category, price, currency="BYN", at="2026-08-13 10:00:00", place=None, qty=1, id=1):
    return {
        "id": id,
        "product_name": product,
        "category_name": category,
        "price": price,
        "currency": currency,
        "purchased_at": at,
        "place": place,
        "qty": qty,
        "unit": "шт",
        "unit_price": price,
        "source": "text",
    }


def test_build_charts_one_per_currency():
    rows = [
        row("молоко", "молочка", 3.0, "BYN"),
        row("наушники", "техника", 10.0, "USD"),
    ]
    charts = reports.build_charts(rows)
    assert [c[0] for c in charts] == ["BYN", "USD"]
    for _, png in charts:
        assert png[:8] == b"\x89PNG\r\n\x1a\n"


def test_report_xlsx_groups_items_under_category():
    import io

    from openpyxl import load_workbook

    rows = [
        row("молоко", "молочка", 1.92),
        row("масло", "молочка", 4.65),
        row("наушники", "техника", 30.0, "USD"),
    ]
    data = reports.build_report_xlsx(rows)
    assert data[:2] == b"PK"
    ws = load_workbook(io.BytesIO(data)).active
    assert ws.sheet_properties.outlinePr.summaryBelow is False
    col_a = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
    assert "молочка" in col_a and "техника" in col_a  # category rows
    assert "молоко" in col_a and "масло" in col_a  # product rows
    levels = [ws.row_dimensions[r].outline_level for r in range(1, ws.max_row + 1)]
    assert 1 in levels and 2 in levels  # product and purchase grouping


def test_expenses_csv_has_bom_and_header():
    data = reports.expenses_csv([row("молоко", "молочка", 1.92, place="Корона")])
    assert data[:3] == b"\xef\xbb\xbf"
    text = data.decode("utf-8-sig")
    assert "категория" in text.splitlines()[0]
    assert "молоко" in text
    assert "Корона" in text
